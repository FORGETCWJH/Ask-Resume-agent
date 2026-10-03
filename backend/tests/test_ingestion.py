from io import BytesIO
from zipfile import ZipInfo

from docx import Document

from app.services.ingestion import TEXT_EXTENSIONS, _chunks, _extract_docx, _zip_member_allowed


def test_chunks_preserve_all_text_with_overlap():
    chunks = _chunks("a" * 3000, size=1000, overlap=100)
    assert len(chunks) == 4
    assert chunks[0][-100:] == chunks[1][:100]


def test_zip_member_boundary_rejects_traversal_and_binary_files():
    assert _zip_member_allowed("src/main.py", ZipInfo("src/main.py"))
    assert not _zip_member_allowed("../secret.env", ZipInfo("../secret.env"))
    assert not _zip_member_allowed("dist/app.exe", ZipInfo("dist/app.exe"))
    assert ".py" in TEXT_EXTENSIONS


def test_zip_member_boundary_rejects_symlink_and_extreme_compression_ratio():
    symlink = ZipInfo("src/link.py")
    symlink.external_attr = (0o120777 << 16)
    compressed = ZipInfo("src/large.py")
    compressed.file_size = 101
    compressed.compress_size = 1

    assert not _zip_member_allowed(symlink.filename, symlink)
    assert not _zip_member_allowed(compressed.filename, compressed)


def test_docx_parser_keeps_table_resume_content():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技能掌握"
    table.cell(0, 1).text = "Python、FastAPI"
    output = BytesIO()
    document.save(output)
    path = __import__("pathlib").Path("data/test-table-resume.docx")
    path.write_bytes(output.getvalue())
    try:
        assert "技能掌握 | Python、FastAPI" in _extract_docx(path)
    finally:
        path.unlink(missing_ok=True)
