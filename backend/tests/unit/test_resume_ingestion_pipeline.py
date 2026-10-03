from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document

from app.services.ingestion import _extract_docx, _extract_pdf_pages


def _write_text_pdf(path: Path, text: str) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((40, 60), text)
    document.save(path)
    document.close()


def _write_blank_pdf(path: Path) -> None:
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()


def _write_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("张三")
    document.add_paragraph("技能掌握：Python、FastAPI")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目名称"
    table.cell(0, 1).text = "内容生产平台"
    document.save(path)


def test_text_pdf_returns_page_text_and_page_number(tmp_path: Path):
    path = tmp_path / "resume.pdf"
    _write_text_pdf(path, "Skills: Python, FastAPI")

    pages, warnings = _extract_pdf_pages(path)

    assert pages == [("Skills: Python, FastAPI", 1)]
    assert warnings == []


def test_image_pdf_uses_ocr_when_text_layer_is_empty(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "image-resume.pdf"
    _write_blank_pdf(path)

    class FakeOCR:
        def __call__(self, _image):
            return ([[None, "技能掌握"], [None, "Python、FastAPI"]], None)

    monkeypatch.setattr("rapidocr_onnxruntime.RapidOCR", lambda: FakeOCR())

    pages, warnings = _extract_pdf_pages(path)

    assert pages == [("技能掌握\nPython、FastAPI", 1)]
    assert warnings == []


def test_docx_parser_keeps_paragraphs_and_tables(tmp_path: Path):
    path = tmp_path / "resume.docx"
    _write_docx(path)

    content = _extract_docx(path)

    assert "张三" in content
    assert "技能掌握：Python、FastAPI" in content
    assert "项目名称 | 内容生产平台" in content


def test_empty_pdf_page_returns_explicit_warning(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "empty-resume.pdf"
    _write_blank_pdf(path)

    class EmptyOCR:
        def __call__(self, _image):
            return ([], None)

    monkeypatch.setattr("rapidocr_onnxruntime.RapidOCR", lambda: EmptyOCR())

    pages, warnings = _extract_pdf_pages(path)

    assert pages == []
    assert warnings == [{"code": "OCR_PAGE_EMPTY", "pageNumber": 1, "message": "第 1 页未识别出文本"}]


def test_single_page_ocr_failure_keeps_failure_as_page_warning(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "failed-page.pdf"
    _write_blank_pdf(path)

    class FailedOCR:
        def __call__(self, _image):
            raise RuntimeError("ocr unavailable")

    monkeypatch.setattr("rapidocr_onnxruntime.RapidOCR", lambda: FailedOCR())

    pages, warnings = _extract_pdf_pages(path)

    assert pages == []
    assert warnings[0]["code"] == "OCR_PAGE_FAILED"
    assert warnings[0]["pageNumber"] == 1


def test_corrupted_pdf_returns_parse_warning_instead_of_raising(tmp_path: Path):
    path = tmp_path / "corrupted.pdf"
    path.write_bytes(b"not a pdf")

    pages, warnings = _extract_pdf_pages(path)

    assert pages == []
    assert warnings[0]["code"] == "PDF_PARSE_FAILED"
