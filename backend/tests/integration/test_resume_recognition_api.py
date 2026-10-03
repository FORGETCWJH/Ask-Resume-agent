from io import BytesIO
import asyncio
import fitz
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.db import SessionLocal
from app.models import Material, ResumeExtractionDraft
from app.config import settings
from app.service.resume_recognition_service import ResumeRecognitionService
from app.services.ingestion import ingest_resume


def _resume_bytes() -> bytes:
    document = Document()
    document.add_heading("技能掌握", level=1)
    document.add_paragraph("后端开发：熟悉 Python、FastAPI、SQLAlchemy。")
    document.add_heading("工作/实习经历", level=1)
    document.add_paragraph("2025年06月-2025年08月 上海示例科技有限公司 AI应用开发实习生")
    document.add_paragraph("项目名称：内容生产平台；核心技术：Python、MySQL。")
    document.add_heading("项目经验", level=1)
    document.add_paragraph("本地生活服务平台；角色：后端开发；核心技术：Redis、RabbitMQ。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _two_page_resume_pdf() -> bytes:
    document = fitz.open()
    first = document.new_page()
    first.insert_text((40, 60), "Resume\nSkills: Python, FastAPI")
    second = document.new_page()
    second.insert_text((40, 60), "Project Experience\nProject: Content Platform")
    output = BytesIO()
    document.save(output)
    document.close()
    return output.getvalue()


def test_recognition_draft_can_be_edited_confirmed_and_used_as_evidence():
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "识别集成测试"})
        assert created.status_code == 201
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert uploaded.status_code == 202
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")

            recognition = client.get(f"/api/v1/materials/{material_id}/recognition")
            assert recognition.status_code == 200
            body = recognition.json()
            assert body["recognitionStatus"] == "awaitingConfirmation"
            assert set(body["draft"]) == {"personalInfo", "skills", "workExperiences", "projects"}
            assert isinstance(body["draft"]["skills"], dict)
            assert "content" in body["draft"]["skills"]
            assert "projects" not in body["draft"]["workExperiences"][0]
            assert "description" in body["draft"]["workExperiences"][0]
            assert "metrics" not in body["draft"]["projects"][0]

            before_confirm = client.get(f"/api/v1/material-sets/{set_id}/evidence")
            assert before_confirm.status_code == 200
            assert not any(item.get("sourcePath", "").startswith("resume-recognition:") for item in before_confirm.json())

            edited = client.patch(
                f"/api/v1/materials/{material_id}/recognition",
                json={"draft": {**body["draft"], "personalInfo": {"targetRole": "Python 后端工程师"}}},
            )
            assert edited.status_code == 200
            assert edited.json()["draft"]["personalInfo"]["targetRole"] == "Python 后端工程师"

            confirmed = client.post(f"/api/v1/materials/{material_id}/recognition/confirm", json={"section": "skills"})
            assert confirmed.status_code == 200
            assert confirmed.json()["confirmedSections"] == ["skills"]
            assert "skills.content" in confirmed.json()["evidenceLinks"]

            evidence = client.get(f"/api/v1/material-sets/{set_id}/evidence")
            assert evidence.status_code == 200
            assert any(item.get("sourcePath") == "resume-recognition:skills" for item in evidence.json())
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204
            with SessionLocal() as db:
                assert db.query(ResumeExtractionDraft).filter(ResumeExtractionDraft.material_id == material_id).count() == 0


def test_new_material_revision_copies_existing_resume_draft_without_reprocessing():
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "版本隔离测试"})
        assert created.status_code == 201
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            resume_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            assert client.post(f"/api/v1/materials/{resume_id}/recognition/confirm", json={"section": "skills"}).status_code == 200

            project = client.post(f"/api/v1/material-sets/{set_id}/materials?kind=projectArchive", files={"file": ("project.zip", b"PK\x05\x06" + b"\x00" * 18, "application/zip")})
            assert project.status_code == 202
            copied_resume_id = next(item["id"] for item in project.json()["materials"] if item["kind"] == "resume")
            copied = client.get(f"/api/v1/materials/{copied_resume_id}/recognition")
            assert copied.status_code == 200
            assert copied.json()["confirmedSections"] == ["skills"]
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_editing_confirmed_section_revokes_its_fact_evidence():
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "确认回收测试"})
        assert created.status_code == 201
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            initial = client.get(f"/api/v1/materials/{material_id}/recognition").json()
            draft = initial["draft"]
            draft["skills"] = {"content": "后端：FastAPI", "evidenceIds": [], "edited": True}
            assert client.patch(f"/api/v1/materials/{material_id}/recognition", json={"draft": draft}).status_code == 200
            assert client.post(f"/api/v1/materials/{material_id}/recognition/confirm", json={"section": "skills"}).status_code == 200
            assert any(item.get("sourcePath") == "resume-recognition:skills" for item in client.get(f"/api/v1/material-sets/{set_id}/evidence").json())

            draft["skills"]["content"] = "FastAPI 和 Redis"
            edited = client.patch(f"/api/v1/materials/{material_id}/recognition", json={"draft": draft})
            assert edited.status_code == 200
            assert "skills" not in edited.json()["confirmedSections"]
            assert not any(item.get("sourcePath") == "resume-recognition:skills" for item in client.get(f"/api/v1/material-sets/{set_id}/evidence").json())
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_configured_llm_does_not_block_resume_upload(monkeypatch):
    scheduled: list[str] = []
    monkeypatch.setattr(settings, "llm_api_key", "configured-for-test")
    monkeypatch.setattr("app.service.material_service.schedule_resume_recognition", scheduled.append)
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "后台识别上传测试"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert uploaded.status_code == 202
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            with SessionLocal() as db:
                draft = db.query(ResumeExtractionDraft).filter(ResumeExtractionDraft.material_id == material_id).one()
                assert draft.status == "pending"
            assert scheduled == [material_id]
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_configured_llm_recognition_polling_can_resume_pending_job(monkeypatch):
    monkeypatch.setattr(settings, "llm_api_key", "configured-for-test")
    # Keep the upload pending so the GET endpoint must schedule the recovery task.
    monkeypatch.setattr("app.service.material_service.schedule_resume_recognition", lambda _material_id: None)

    async def fake_background_job(_material_id: str) -> None:
        return None

    monkeypatch.setattr("app.service.resume_recognition_service._run_resume_recognition", fake_background_job)
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "识别轮询恢复测试"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            polled = client.get(f"/api/v1/materials/{material_id}/recognition")
            assert polled.status_code == 200
            assert polled.json()["recognitionStatus"] == "processing"
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_local_recognition_never_calls_llm_even_when_provider_is_configured(monkeypatch):
    monkeypatch.setattr(settings, "llm_api_key", "configured-for-test")

    async def fail_if_called(*_args, **_kwargs):
        raise AssertionError("local recognition must not call LLM")

    monkeypatch.setattr("app.infrastructure.llm.request_resume_extraction", fail_if_called)
    monkeypatch.setattr("app.service.material_service.schedule_resume_recognition", lambda _material_id: None)
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "本地识别不调用模型"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            with SessionLocal() as db:
                material = db.get(Material, material_id)
                ingest_resume(db, material, Path(material.storage_path))
                asyncio.run(ResumeRecognitionService(db).process(material_id))
                draft = db.query(ResumeExtractionDraft).filter(ResumeExtractionDraft.material_id == material_id).one()
                assert draft.status == "awaitingConfirmation"
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_upload_starts_local_recognition_without_llm_configuration(monkeypatch):
    monkeypatch.setattr(settings, "llm_api_key", "")
    scheduled: list[str] = []
    monkeypatch.setattr("app.service.material_service.schedule_resume_recognition", scheduled.append)
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "本地上传任务测试"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert uploaded.status_code == 202
            material = next(item for item in uploaded.json()["materials"] if item["kind"] == "resume")
            assert material["status"] == "processing"
            assert scheduled == [material["id"]]
            with SessionLocal() as db:
                draft = db.query(ResumeExtractionDraft).filter(ResumeExtractionDraft.material_id == material["id"]).one()
                assert draft.status == "pending"
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_recognition_can_confirm_all_sections():
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "全部确认测试"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.docx", _resume_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            confirmed = client.post(f"/api/v1/materials/{material_id}/recognition/confirm", json={"section": "all"})
            assert confirmed.status_code == 200
            assert confirmed.json()["confirmedSections"] == ["personalInfo", "skills", "workExperiences", "projects"]
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204


def test_pdf_page_retry_keeps_other_page_evidence():
    with TestClient(app) as client:
        created = client.post("/api/v1/material-sets", json={"title": "页级重试测试"})
        set_id = created.json()["id"]
        try:
            uploaded = client.post(
                f"/api/v1/material-sets/{set_id}/materials?kind=resume",
                files={"file": ("resume.pdf", _two_page_resume_pdf(), "application/pdf")},
            )
            material_id = next(item["id"] for item in uploaded.json()["materials"] if item["kind"] == "resume")
            before = client.get(f"/api/v1/material-sets/{set_id}/evidence").json()
            assert {item["pageNumber"] for item in before if item["pageNumber"]} == {1, 2}
            retried = client.post(
                f"/api/v1/materials/{material_id}/recognition/retry",
                json={"scope": "ocr", "pageNumber": 1},
            )
            assert retried.status_code == 200
            after = client.get(f"/api/v1/material-sets/{set_id}/evidence").json()
            assert any(item["pageNumber"] == 2 and "Project Experience" in item["content"] for item in after)
        finally:
            assert client.delete(f"/api/v1/material-sets/{set_id}").status_code == 204
